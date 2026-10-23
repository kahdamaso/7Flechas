"""Importação dos cadastros com uma planilha FICTÍCIA no mesmo formato da real."""

import openpyxl
import pytest

from app.importacao import importar
from app.models import Entidade, HistoricoCategoria, Parentesco, Pessoa
from app.utils import normalizar_cpf


@pytest.fixture
def planilha(tmp_path):
    livro = openpyxl.Workbook()
    livro.remove(livro.active)

    presenca = livro.create_sheet("PRESENÇA")
    presenca.append(["MEDIUNS", "CATOGORIA", "PRESENTE"])
    for linha in [
        ["Igor Teste Silva", "Corrente", "SIM"],
        ["Káh Exemplo Souza", "Corrente", "NÃO"],
        ["Thiago Soares Costa", "Corrente", "SIM"],
        ["Thiago Soares Cecilio", "Corrente", "SIM"],
        ["Railany Ficticia Vieira", "Corrente", "SIM"],
        ["Maria Nunes Campos Ribeiro", "Corrente", "SIM"],
        ["Mylene Teste Melo", "Corrente", "SIM"],
        ["Ana Auxiliar", "Auxiliar", "NÃO"],
        ["Bia Colaboradora", "Colaborador", "SIM"],
    ]:
        presenca.append(linha)

    retorno = livro.create_sheet("RETORNO MEDIUNS")
    retorno.append(["MEDIUNS", "Tipo", "RETORNO"])
    retorno.append(["Ana Auxiliar", "Corrente", ""])  # diverge da PRESENÇA
    retorno.append(["Novo Membro", "Auxiliar", ""])

    ficha = livro.create_sheet("Cópia de GUIA")
    ficha.append(["Mediun", "Caboclo", "Preto Velhos", "Exu", "Pombo Gira", "Cangaço"])
    ficha.append(["Igor Teste Silva", "7 Flechas", "Pai João", "7 catacumbas", None, "Cangaceiro"])
    # mesma pessoa com as palavras do nome em outra ordem
    ficha.append(["Thiago Costa Soares", "Sutão da Matas", None, "Ze Caveira / Tiriri", None, None])
    ficha.append(["Pessoa Sem Categoria", "Jurema", None, None, None, None])

    guia = livro.create_sheet("GUIA")
    guia.append(["Caboclo", "Preto Velhos", "Exu"])
    guia.append(["7Flechas", "Pai Jacó", "Tiriri"])  # nomes de entidades: ignorados
    guia.append(["Pai Igor", "Mãe Ralainy", "Thiago DV"])
    guia.append(["Mylene", "GERAL", "Pai Káh"])

    consulente = livro.create_sheet("CONSULENTE")
    consulente.append(["CONSULENTE", "CPF", "PARENTESCO", "CIDADE", "PERTENCE ALGUMA CASA? QUAL",
                       "COMO FICOU SABENDO", "TELEFONE"])
    consulente.append(["JOAO EXEMPLO", "529.982.247-25", "MARIA NUNES", "BO DESPACHO", "NÃO",
                       "INSTAGRAM", 37998887777.0])
    consulente.append(["CARLA FICTICIA", 1234567890.0, "PAI KAH, MILENE", "Nova Serrana", "OASIS",
                       None, None])
    consulente.append(["PEDRO CRIANCA", "CRIANÇA", "THIAGO", "0.0", "SIM", None, None])
    consulente.append(["Igor Teste Silva", None, "NÃO", "Bom Despacho", None, None, None])

    caminho = tmp_path / "TUC7FJ.xlsx"
    livro.save(caminho)
    return caminho


def _pessoa(nome):
    return Pessoa.query.filter_by(nome=nome).one()


def test_normalizar_cpf():
    assert normalizar_cpf("529.982.247-25") == "52998224725"
    assert normalizar_cpf("1234567890") == "01234567890"  # zero à esquerda perdido no Excel
    assert normalizar_cpf("111.111.111-11") is None
    assert normalizar_cpf("CRIANÇA") is None


def test_categorias_e_historico(app, planilha):
    rel = importar(planilha)
    assert _pessoa("Igor Teste Silva").categoria == "Corrente"
    # PRESENÇA prevalece sobre RETORNO MEDIUNS, e a divergência é relatada
    assert _pessoa("Ana Auxiliar").categoria == "Auxiliar"
    assert any("Ana Auxiliar" in i for i in rel.pendencias["Categorias divergentes entre abas (mantida a primeira)"])
    assert _pessoa("Novo Membro").categoria == "Auxiliar"
    assert _pessoa("Bia Colaboradora").categoria == "Colaborador"
    h = HistoricoCategoria.query.filter_by(pessoa_id=_pessoa("Igor Teste Silva").id).one()
    assert h.de is None and h.para == "Corrente"


def test_ficha_espiritual(app, planilha):
    rel = importar(planilha)
    igor = _pessoa("Igor Teste Silva")
    ficha = {(e.linha.nome, e.nome) for e in igor.entidades}
    assert ficha == {("Caboclo", "7 Flechas"), ("Preto Velho", "Pai João"),
                     ("Exu", "7 catacumbas"), ("Cangaço", "Cangaceiro")}
    assert igor.incorpora
    thiago = _pessoa("Thiago Soares Costa")  # 'Thiago Costa Soares' na ficha
    assert {e.nome for e in thiago.entidades} == {"Sultão das Matas", "Zé Caveira", "Tiriri"}
    assert Pessoa.query.filter(Pessoa.nome.like("Thiago%")).count() == 2
    assert _pessoa("Pessoa Sem Categoria").categoria == "Corrente"
    assert "Pessoa Sem Categoria" in rel.pendencias[
        "Médiuns da ficha sem categoria em outra aba (cadastrados como Corrente)"]


def test_apelidos(app, planilha):
    rel = importar(planilha)
    assert _pessoa("Igor Teste Silva").apelido == "Pai Igor"
    assert _pessoa("Káh Exemplo Souza").apelido == "Pai Káh"
    assert _pessoa("Railany Ficticia Vieira").apelido == "Mãe Ralainy"  # grafia parecida
    # só o primeiro nome: marca que incorpora, sem virar apelido
    mylene = _pessoa("Mylene Teste Melo")
    assert mylene.incorpora and mylene.apelido is None
    # dois Thiagos: fica para conferir; 'Pai Jacó' é entidade, não apelido
    pendentes = rel.pendencias["Apelidos não identificados (preencha o apelido no cadastro do médium)"]
    assert pendentes == ["Thiago DV"]


def test_consulentes_e_parentescos(app, planilha):
    rel = importar(planilha)
    joao = _pessoa("Joao Exemplo")
    assert joao.categoria == "Consulente"
    assert joao.cpf == "52998224725" and joao.telefone == "37998887777"
    assert joao.cidade == "Bom Despacho" and joao.outra_casa is None and joao.como_soube == "INSTAGRAM"
    maria = _pessoa("Maria Nunes Campos Ribeiro")
    assert Parentesco.entre(joao.id, maria.id) is not None

    carla = _pessoa("Carla Ficticia")
    assert carla.cpf == "01234567890" and carla.outra_casa == "OASIS"
    assert Parentesco.entre(carla.id, _pessoa("Káh Exemplo Souza").id)
    assert Parentesco.entre(carla.id, _pessoa("Mylene Teste Melo").id)  # 'MILENE'
    assert any("Mylene" in i for i in rel.pendencias["Parentescos vinculados por grafia parecida (confira)"])
    assert carla.parentesco_informado == "PAI KAH, MILENE"

    pedro = _pessoa("Pedro Crianca")
    assert pedro.cpf is None and "CRIANÇA" in pedro.observacoes
    assert pedro.cidade is None and pedro.outra_casa == "Sim (não informou qual)"
    assert any("'thiago'" in i for i in rel.pendencias["Parentescos não identificados (cadastre no prontuário)"])

    # médium que também está na aba CONSULENTE não vira um segundo cadastro
    assert Pessoa.query.filter_by(nome="Igor Teste Silva").count() == 1
    assert _pessoa("Igor Teste Silva").categoria == "Corrente"


def test_reimportacao_nao_duplica(app, planilha):
    importar(planilha)
    totais = (Pessoa.query.count(), Entidade.query.count(), Parentesco.query.count(),
              HistoricoCategoria.query.count())
    importar(planilha)
    assert totais == (Pessoa.query.count(), Entidade.query.count(), Parentesco.query.count(),
                      HistoricoCategoria.query.count())


def test_reimportacao_preserva_o_que_foi_digitado(app, planilha):
    from app.extensions import db

    importar(planilha)
    joao = _pessoa("Joao Exemplo")
    joao.cidade = "Divinópolis"
    joao.mudar_categoria("Colaborador")
    db.session.commit()
    importar(planilha)
    joao = _pessoa("Joao Exemplo")
    assert joao.cidade == "Divinópolis" and joao.categoria == "Colaborador"
